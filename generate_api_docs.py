#!/usr/bin/env python3
"""generate_api_docs.py

Generates simple, human-readable API documentation — as Markdown or as a
Word (.docx) document — from one or more service_catalog.json files
produced by ace_service_catalog.py.

For every service it documents: the input (parameters), a synthesized
sample request, the output (response schema), a synthesized sample
response, and the backend operations that service calls.

"Synthesized" is the operative word: a sample value is only produced when
it can be derived from something actually documented in the catalog (a
schema's own "example", an enum's first value, or a primitive type's
placeholder). Where a schema is a bare reference (e.g. "#/definitions/Profile")
with no inline body — which is how ace_service_catalog.py deliberately
preserves undocumented-here schemas rather than guessing their shape — this
tool says so explicitly instead of inventing fields that were never in the
catalog.

Usage
-----
    python generate_api_docs.py --catalog service_catalog.json --output api_documentation.md
    python generate_api_docs.py --catalog catalog1.json catalog2.json --output api_documentation.md
    python generate_api_docs.py --catalog service_catalog.json --output api_documentation.docx
    python generate_api_docs.py --catalog service_catalog.json --output out.md --format docx
        (--format overrides whatever the --output extension would imply)

The .docx output requires python-docx (pip install python-docx); Markdown
output has no extra dependency.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlencode


PLACEHOLDER_BY_TYPE = {
    "string": "string",
    "integer": 0,
    "number": 0,
    "boolean": False,
}


# ============================================================================
# Schema / sample-value synthesis (shared by both renderers)
# ============================================================================

def synthesize_example(schema: Any, depth: int = 0) -> Any:
    """Best-effort example value from a JSON schema fragment.

    Returns None when nothing can be derived (a bare $ref string, or an
    empty/unknown schema) — callers must render that as "schema reference
    only", never fabricate a body.
    """
    if schema is None or isinstance(schema, str) or not isinstance(schema, dict):
        return None
    if depth > 8:
        return None
    if "example" in schema:
        return schema["example"]

    schema_type = schema.get("type")
    if schema_type == "object" or "properties" in schema:
        out: Dict[str, Any] = {}
        for name, sub in (schema.get("properties") or {}).items():
            val = synthesize_example(sub, depth + 1)
            if val is None and isinstance(sub, dict):
                if sub.get("enum"):
                    val = sub["enum"][0]
                elif sub.get("type"):
                    val = PLACEHOLDER_BY_TYPE.get(sub["type"])
            if val is not None:
                out[name] = val
        return out or None

    if schema_type == "array":
        item = synthesize_example(schema.get("items"), depth + 1)
        return [item] if item is not None else []

    if schema.get("enum"):
        return schema["enum"][0]
    if schema_type in PLACEHOLDER_BY_TYPE:
        return PLACEHOLDER_BY_TYPE[schema_type]
    return None


def schema_reference_note(schema: Any) -> Optional[str]:
    """If `schema` is a bare $ref string, return an explanatory note; else None."""
    if isinstance(schema, str):
        return f"Conforms to schema {schema} (definition not embedded in this catalog)."
    return None


def fill_path_placeholders(path: str, params: List[Dict[str, Any]]) -> str:
    """Fill {name} path tokens only where the catalog actually documents a
    real value (an example, enum, or default) — otherwise leave the token
    as-is. A synthesized literal like "string" or "0" reads as a real
    sample value when it isn't one; the original {customer_id} token is
    self-evidently a placeholder and doesn't misrepresent anything.
    """
    filled = path
    for p in params:
        if p.get("in") != "path":
            continue
        name = p.get("name")
        if not name:
            continue
        example = p.get("example")
        if example is None and p.get("enum"):
            example = p["enum"][0]
        if example is None:
            example = p.get("default")
        if example is None:
            continue
        filled = filled.replace("{" + name + "}", str(example))
    return filled


def build_sample_request(op_input: Dict[str, Any]) -> str:
    method = op_input.get("method", "GET")
    params = op_input.get("parameters") or []
    path = fill_path_placeholders(op_input.get("full_path") or op_input.get("path") or "", params)

    query_params = {}
    for p in params:
        if p.get("in") != "query":
            continue
        example = p.get("example")
        if example is None and p.get("enum"):
            example = p["enum"][0]
        if example is None:
            example = p.get("default")
        if example is None and not p.get("required"):
            continue  # skip optional query params with nothing to show
        if example is None:
            example = f"<{p.get('type','string')}>"
        query_params[p["name"]] = example
    if query_params:
        path = path + "?" + urlencode(query_params)

    header_lines = []
    for p in params:
        if p.get("in") != "header":
            continue
        example = p.get("example")
        if example is None and p.get("enum"):
            example = p["enum"][0]
        if example is None:
            example = p.get("default", f"<{p.get('type','string')}>")
        header_lines.append(f"{p['name']}: {example}")

    body_param = next((p for p in params if p.get("in") == "body"), None)
    body_text = None
    if body_param is not None:
        schema = body_param.get("schema")
        if schema is not None:
            example = synthesize_example(schema)
            if example is not None:
                body_text = json.dumps(example, indent=2)
        # a body param that's a pure $ref (schema_ref) has no inline shape to
        # synthesize from — left unset deliberately, not fabricated.

    lines = [f"{method} {path} HTTP/1.1"]
    lines.extend(header_lines)
    request_text = "\n".join(lines)
    if body_text:
        request_text += "\n\n" + body_text
    return request_text


# ============================================================================
# Shared row-data builders — one source of truth consumed by both the
# Markdown and the .docx renderer, so the two output formats can never
# silently drift apart.
# ============================================================================

def parameter_rows(params: List[Dict[str, Any]]) -> List[Tuple[str, str, str, str, str]]:
    rows = []
    for p in params:
        desc_bits = []
        if p.get("description"):
            desc_bits.append(p["description"])
        if p.get("enum"):
            desc_bits.append("enum: " + ", ".join(str(v) for v in p["enum"]))
        if p.get("schema_ref"):
            desc_bits.append(f"ref: {p['schema_ref']}")
        elif p.get("schema") is not None:
            desc_bits.append("inline object — see sample request body")
        desc = " — ".join(desc_bits)
        type_cell = p.get("type") or ("object" if p.get("schema") else "")
        rows.append((p.get("name", ""), p.get("in", ""), type_cell, "Yes" if p.get("required") else "No", desc))
    return rows


def backend_rows(backends: List[Dict[str, Any]]) -> List[Tuple[str, str, str, str, str, str]]:
    rows = []
    for b in backends:
        if b.get("url"):
            url_text = b["url"]
        elif b.get("configured_url"):
            url_text = f"dynamic — configured as {b['configured_url']}"
        else:
            url_text = "unresolved"
        rows.append((
            b.get("node_name", ""), b.get("backend_system", ""), b.get("protocol", ""),
            url_text, b.get("resolution", ""), b.get("confidence", ""),
        ))
    return rows


def response_blocks(response_schemas: Dict[str, Any]) -> List[Tuple[str, Optional[str], Optional[str]]]:
    """Returns [(status_code, ref_note_or_None, json_text_or_None), ...]."""
    blocks = []
    for code, schema in (response_schemas or {}).items():
        ref_note = schema_reference_note(schema)
        if ref_note:
            blocks.append((str(code), ref_note, None))
        else:
            example = synthesize_example(schema)
            blocks.append((str(code), None, json.dumps(example if example is not None else schema, indent=2)))
    return blocks


def sample_response(response_schemas: Dict[str, Any]) -> Tuple[Optional[str], Optional[str]]:
    """Returns (ref_note_or_None, json_text_or_None) for the first 2xx response, or (None, None) if there isn't one."""
    resp = response_schemas or {}
    success_code = next((c for c in resp if str(c).startswith("2")), None)
    if success_code is None:
        return None, None
    schema = resp[success_code]
    ref_note = schema_reference_note(schema)
    if ref_note:
        return ref_note, None
    example = synthesize_example(schema)
    return None, json.dumps(example if example is not None else schema, indent=2)


# ============================================================================
# Markdown renderer
# ============================================================================

def md_table(headers: List[str], rows: List[Tuple], empty_message: str) -> str:
    if not rows:
        return f"_{empty_message}_\n"
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for row in rows:
        lines.append("| " + " | ".join(str(c).replace("|", "\\|").replace("\n", " ") for c in row) + " |")
    return "\n".join(lines) + "\n"


def render_service_md(svc: Dict[str, Any], catalog_name: str) -> str:
    name = svc.get("service_name", "(unnamed)")
    op_input = svc.get("input", {}) or {}
    swagger = svc.get("swagger", {}) or {}
    backends = svc.get("backends", []) or []
    method = op_input.get("method", "GET")
    full_path = op_input.get("full_path") or op_input.get("path") or ""

    out = [f"## {name}", "", f"**{method}** `{full_path}`", ""]
    if swagger.get("summary"):
        out += [f"**{swagger['summary']}**", ""]
    if swagger.get("description"):
        out += [swagger["description"], ""]
    tags = swagger.get("tags") or []
    if tags:
        out += [f"_Tags: {', '.join(tags)}_", ""]
    out += [f"_Source catalog: {catalog_name}_", ""]

    out += ["### Input", "", md_table(
        ["Name", "In", "Type", "Required", "Description"],
        parameter_rows(op_input.get("parameters") or []), "No parameters",
    )]

    out += ["### Sample Request", "", "```http", build_sample_request(op_input), "```", ""]

    out += ["### Output", ""]
    blocks = response_blocks(swagger.get("response_schemas") or {})
    if not blocks:
        out += ["_No response schemas documented for this operation._", ""]
    for code, ref_note, json_text in blocks:
        out += [f"**HTTP {code}**", ""]
        out += [f"_{ref_note}_", ""] if ref_note else ["```json", json_text, "```", ""]

    out += ["### Sample Response", ""]
    ref_note, json_text = sample_response(swagger.get("response_schemas") or {})
    if ref_note:
        out += [f"_{ref_note}_"]
    elif json_text:
        out += ["```json", json_text, "```"]
    else:
        out += ["_No 2xx response documented for this operation._"]
    out.append("")

    out += ["### Backend Operations Used", "", md_table(
        ["Node", "Backend System", "Protocol", "URL", "Resolution", "Confidence"],
        backend_rows(backends), "This service does not call any backend system directly or via downstream flows",
    )]

    out += ["---", ""]
    return "\n".join(out)


def render_catalog_md(catalog: Dict[str, Any]) -> str:
    catalog_name = catalog.get("catalog_name", "Unnamed Catalog")
    services = catalog.get("services", []) or []
    out = [f"# {catalog_name}", ""]
    if catalog.get("catalog_version"):
        out += [f"Version: {catalog['catalog_version']}", ""]
    out += [f"{len(services)} service(s) documented below.", "", "## Contents", ""]
    for svc in services:
        anchor = svc.get("service_name", "unnamed").lower().replace(" ", "-")
        out.append(f"- [{svc.get('service_name','(unnamed)')}](#{anchor})")
    out.append("")
    for svc in services:
        out.append(render_service_md(svc, catalog_name))
    return "\n".join(out)


def write_markdown(catalogs: List[Dict[str, Any]], output_path: Path) -> None:
    output_path.write_text("\n\n".join(render_catalog_md(c) for c in catalogs), encoding="utf-8")


# ============================================================================
# .docx renderer
# ============================================================================

def write_docx(catalogs: List[Dict[str, Any]], output_path: Path) -> None:
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise RuntimeError(
            "Writing a .docx file requires python-docx, which is not installed.\n"
            "Install it with: pip install python-docx"
        ) from exc

    doc = Document()

    def add_code_block(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        for i, line in enumerate(text.split("\n")):
            if i > 0:
                p.add_run().add_break()
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(9)

    def add_table(headers: List[str], rows: List[Tuple], empty_message: str) -> None:
        if not rows:
            italic = doc.add_paragraph()
            r = italic.add_run(empty_message)
            r.italic = True
            return
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(h)
            r.bold = True
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val) if val not in (None, "") else ""
        doc.add_paragraph()

    def add_response_blocks(blocks: List[Tuple[str, Optional[str], Optional[str]]]) -> None:
        if not blocks:
            doc.add_paragraph("No response schemas documented for this operation.")
            return
        for code, ref_note, json_text in blocks:
            p = doc.add_paragraph()
            p.add_run(f"HTTP {code}").bold = True
            if ref_note:
                r = doc.add_paragraph().add_run(ref_note)
                r.italic = True
            else:
                add_code_block(json_text or "")

    for catalog in catalogs:
        catalog_name = catalog.get("catalog_name", "Unnamed Catalog")
        services = catalog.get("services", []) or []

        doc.add_heading(catalog_name, level=1)
        if catalog.get("catalog_version"):
            doc.add_paragraph(f"Version: {catalog['catalog_version']}")
        doc.add_paragraph(f"{len(services)} service(s) documented below.")

        for svc in services:
            name = svc.get("service_name", "(unnamed)")
            op_input = svc.get("input", {}) or {}
            swagger = svc.get("swagger", {}) or {}
            backends = svc.get("backends", []) or []
            method = op_input.get("method", "GET")
            full_path = op_input.get("full_path") or op_input.get("path") or ""

            doc.add_heading(name, level=2)
            p = doc.add_paragraph()
            p.add_run(f"{method} ").bold = True
            r = p.add_run(full_path)
            r.font.name = "Consolas"

            if swagger.get("summary"):
                r = doc.add_paragraph().add_run(swagger["summary"])
                r.bold = True
            if swagger.get("description"):
                doc.add_paragraph(swagger["description"])
            tags = swagger.get("tags") or []
            if tags:
                r = doc.add_paragraph().add_run(f"Tags: {', '.join(tags)}")
                r.italic = True
            r = doc.add_paragraph().add_run(f"Source catalog: {catalog_name}")
            r.italic = True

            doc.add_heading("Input", level=3)
            add_table(
                ["Name", "In", "Type", "Required", "Description"],
                parameter_rows(op_input.get("parameters") or []), "No parameters.",
            )

            doc.add_heading("Sample Request", level=3)
            add_code_block(build_sample_request(op_input))

            doc.add_heading("Output", level=3)
            add_response_blocks(response_blocks(swagger.get("response_schemas") or {}))

            doc.add_heading("Sample Response", level=3)
            ref_note, json_text = sample_response(swagger.get("response_schemas") or {})
            if ref_note:
                r = doc.add_paragraph().add_run(ref_note)
                r.italic = True
            elif json_text:
                add_code_block(json_text)
            else:
                doc.add_paragraph("No 2xx response documented for this operation.")

            doc.add_heading("Backend Operations Used", level=3)
            add_table(
                ["Node", "Backend System", "Protocol", "URL", "Resolution", "Confidence"],
                backend_rows(backends), "This service does not call any backend system directly or via downstream flows.",
            )

            sep = doc.add_paragraph()
            sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sep.add_run("* * *")

    doc.save(str(output_path))


# ============================================================================
# CLI
# ============================================================================

def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate simple per-service API documentation (Markdown or .docx) from service_catalog.json file(s).",
    )
    parser.add_argument("--catalog", required=True, nargs="+", help="One or more service_catalog.json files.")
    parser.add_argument("--output", required=True, help="Output file path (.md or .docx).")
    parser.add_argument("--format", choices=["markdown", "docx"], default=None,
                         help="Output format. Defaults to whatever --output's extension implies "
                              "(.docx -> docx, anything else -> markdown).")
    args = parser.parse_args(argv if argv is not None else sys.argv[1:])

    catalogs = []
    for catalog_path in args.catalog:
        path = Path(catalog_path)
        if not path.is_file():
            print(f"ERROR: catalog file not found: {catalog_path}", file=sys.stderr)
            return 1
        with path.open("r", encoding="utf-8") as f:
            catalog = json.load(f)
        if not isinstance(catalog.get("services"), list):
            print(f"ERROR: {catalog_path} does not contain a \"services\" array", file=sys.stderr)
            return 1
        catalogs.append(catalog)

    output_path = Path(args.output)
    fmt = args.format or ("docx" if output_path.suffix.lower() == ".docx" else "markdown")

    try:
        if fmt == "docx":
            write_docx(catalogs, output_path)
        else:
            write_markdown(catalogs, output_path)
    except RuntimeError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    total_services = sum(len(c.get("services", [])) for c in catalogs)
    print(f"Wrote {output_path} ({fmt}, {total_services} service(s) across {len(catalogs)} catalog(s))")
    return 0


if __name__ == "__main__":
    sys.exit(main())
